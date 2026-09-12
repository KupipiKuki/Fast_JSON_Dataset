#!/usr/bin/env python3
"""Build deterministic, sharded JSON lookup datasets from common documents."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
import unicodedata
from pathlib import Path
from typing import Any, Iterable

SCHEMA_VERSION = "1.0.0"
SUPPORTED = {".xlsx", ".xlsm", ".csv", ".docx", ".pdf"}


def canonical_json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def normalize_key(value: Any) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value))
    return re.sub(r"\s+", " ", text).strip().casefold()


def clean_value(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int, float)):
        return value
    return unicodedata.normalize("NFKC", str(value)).strip()


def safe_header(value: Any, number: int) -> str:
    text = clean_value(value)
    text = str(text) if text not in (None, "") else f"column_{number}"
    text = re.sub(r"[^0-9A-Za-z_]+", "_", text).strip("_").lower()
    return text or f"column_{number}"


def unique_headers(values: Iterable[Any]) -> list[str]:
    result, seen = [], {}
    for i, value in enumerate(values, 1):
        base = safe_header(value, i)
        seen[base] = seen.get(base, 0) + 1
        result.append(base if seen[base] == 1 else f"{base}_{seen[base]}")
    return result


def parse_number_spec(spec: str | None, maximum: int) -> list[int]:
    if not spec or spec == "*":
        return list(range(1, maximum + 1))
    values: set[int] = set()
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            start, end = (int(x) for x in part.split("-", 1))
            values.update(range(start, end + 1))
        else:
            values.add(int(part))
    invalid = [n for n in values if n < 1 or n > maximum]
    if invalid:
        raise ValueError(f"Selection outside 1..{maximum}: {invalid}")
    return sorted(values)


def make_record(namespace: str, source_name: str, source_type: str, locator: dict[str, Any],
                data: dict[str, Any], id_fields: list[str]) -> dict[str, Any]:
    usable = id_fields and all(data.get(k) not in (None, "") for k in id_fields)
    identity = {"namespace": namespace, "source": source_name, "locator_scope": {
        k: v for k, v in locator.items() if k not in {"row", "line", "paragraph"}
    }}
    if usable:
        identity["keys"] = {k: normalize_key(data[k]) for k in id_fields}
    else:
        identity["locator"] = locator
    record_id = hashlib.sha256(canonical_json(identity).encode()).hexdigest()[:24]
    content_hash = hashlib.sha256(canonical_json(data).encode()).hexdigest()
    return {"record_id": record_id, "content_hash": content_hash,
            "source": {"name": source_name, "type": source_type},
            "locator": locator, "data": data}


def row_records(rows: list[list[Any]], header_row: int, namespace: str, source_name: str,
                source_type: str, locator_base: dict[str, Any], id_fields: list[str],
                absolute_start_row: int = 1) -> list[dict[str, Any]]:
    if not rows:
        return []
    hi = header_row - 1
    if hi < 0 or hi >= len(rows):
        raise ValueError(f"header_row {header_row} is outside selected content")
    headers = unique_headers(rows[hi])
    records = []
    for rel, row in enumerate(rows[hi + 1:], hi + 2):
        values = list(row) + [None] * (len(headers) - len(row))
        data = {headers[i]: clean_value(values[i]) for i in range(len(headers))}
        if all(v in (None, "") for v in data.values()):
            continue
        locator = dict(locator_base)
        locator["row"] = absolute_start_row + rel - 1
        records.append(make_record(namespace, source_name, source_type, locator, data, id_fields))
    return records


def excel_records(path: Path, cfg: dict[str, Any], args: argparse.Namespace, audit: dict[str, Any]) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.cell import range_boundaries
    except ImportError as e:
        raise RuntimeError("Excel extraction requires openpyxl") from e
    options = cfg.get("options", {}).get("excel", {})
    wb = load_workbook(path, read_only=True, data_only=options.get("data_only", True))
    configured = cfg.get("selection", {}).get("spreadsheets", {}).get("sheets", [])
    selected: list[dict[str, Any]] = []
    if args.ranges:
        for value in args.ranges:
            if "!" not in value:
                raise ValueError(f"Excel range must be Sheet!A1:D10: {value}")
            name, cell_range = value.split("!", 1)
            selected.append({"name": name, "range": cell_range, "header_row": 1})
    elif args.sheets:
        selected = [{"name": name, "range": None, "header_row": 1} for name in args.sheets]
    else:
        selected = configured or [{"name": name, "range": None, "header_row": 1} for name in wb.sheetnames]
    output = []
    for item in selected:
        name = item["name"] if isinstance(item, dict) else str(item)
        if name not in wb.sheetnames:
            audit["errors"].append(f"Worksheet not found: {name}")
            continue
        ws = wb[name]
        cell_range = item.get("range") if isinstance(item, dict) else None
        header_row = int(item.get("header_row", 1)) if isinstance(item, dict) else 1
        if cell_range:
            min_col, min_row, max_col, max_row = range_boundaries(cell_range)
        else:
            min_col, min_row, max_col, max_row = 1, 1, ws.max_column, ws.max_row
            cell_range = f"A1:{ws.cell(max_row, max_col).coordinate}"
        rows = [[cell.value for cell in row] for row in ws.iter_rows(
            min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col)]
        recs = row_records(rows, header_row, cfg["namespace"], path.name, "excel",
                           {"sheet": name, "range": cell_range}, cfg.get("record_id_fields", []), min_row)
        output.extend(recs)
        audit["processed_selections"].append({"sheet": name, "range": cell_range, "records": len(recs)})
    wb.close()
    return output


def csv_records(path: Path, cfg: dict[str, Any], audit: dict[str, Any]) -> list[dict[str, Any]]:
    spec = cfg.get("selection", {}).get("csv", {})
    delimiter = cfg.get("options", {}).get("csv", {}).get("delimiter", ",")
    encoding = cfg.get("options", {}).get("csv", {}).get("encoding", "utf-8-sig")
    with path.open("r", encoding=encoding, newline="") as f:
        rows = [list(r) for r in csv.reader(f, delimiter=delimiter)]
    start_row = 1
    if spec.get("range"):
        try:
            from openpyxl.utils.cell import range_boundaries
        except ImportError as e:
            raise RuntimeError("CSV A1 ranges require openpyxl") from e
        min_col, min_row, max_col, max_row = range_boundaries(spec["range"])
        rows = [r[min_col - 1:max_col] for r in rows[min_row - 1:max_row]]
        start_row = min_row
    recs = row_records(rows, int(spec.get("header_row", 1)), cfg["namespace"], path.name, "csv",
                       {"range": spec.get("range", "*")}, cfg.get("record_id_fields", []), start_row)
    audit["processed_selections"].append({"range": spec.get("range", "*"), "records": len(recs)})
    return recs


def word_records(path: Path, cfg: dict[str, Any], args: argparse.Namespace, audit: dict[str, Any]) -> list[dict[str, Any]]:
    try:
        from docx import Document
    except ImportError as e:
        raise RuntimeError("Word extraction requires python-docx") from e
    doc = Document(path)
    spec = cfg.get("selection", {}).get("word", {})
    output = []
    paragraph_spec = spec.get("paragraphs")
    if paragraph_spec:
        for number in parse_number_spec(str(paragraph_spec), len(doc.paragraphs)):
            text = clean_value(doc.paragraphs[number - 1].text)
            if text:
                output.append(make_record(cfg["namespace"], path.name, "word",
                    {"paragraph": number}, {"text": text}, cfg.get("record_id_fields", [])))
        audit["processed_selections"].append({"paragraphs": paragraph_spec,
                                               "records": len(output)})
    table_items = [{"index": i, "header_row": 1} for i in args.word_tables] if args.word_tables else spec.get("tables", [])
    for item in table_items:
        index = int(item["index"] if isinstance(item, dict) else item)
        if index < 0 or index >= len(doc.tables):
            audit["errors"].append(f"Word table not found: {index}")
            continue
        rows = [[cell.text for cell in row.cells] for row in doc.tables[index].rows]
        recs = row_records(rows, int(item.get("header_row", 1)) if isinstance(item, dict) else 1,
                           cfg["namespace"], path.name, "word",
                           {"table": index}, cfg.get("record_id_fields", []))
        output.extend(recs)
        audit["processed_selections"].append({"table": index, "records": len(recs)})
    return output


def pdf_records(path: Path, cfg: dict[str, Any], args: argparse.Namespace, audit: dict[str, Any]) -> list[dict[str, Any]]:
    pdf_cfg = cfg.get("selection", {}).get("pdf", {})
    mode = pdf_cfg.get("mode", "lines")
    spec = args.pages or pdf_cfg.get("pages", "*")
    output = []
    structured_modes = {
        "setting_assignments", "lines_and_setting_assignments",
        "logic_equations", "lines_and_logic_equations",
        "selogic_programming", "lines_and_selogic_programming",
    }
    if mode in structured_modes:
        try:
            import pdfplumber
        except ImportError as e:
            raise RuntimeError(f"PDF {mode} mode requires pdfplumber") from e
        logic_mode = mode in {"logic_equations", "lines_and_logic_equations"}
        programming_mode = mode in {"selogic_programming", "lines_and_selogic_programming"}
        include_lines = mode.startswith("lines_and_")
        if programming_mode:
            # Programming manuals can contain modern := and legacy = examples.
            # Anchor at line start and require uppercase SEL identifiers to reject prose.
            assignment = re.compile(
                r"^\s*(?P<setting_name>[A-Za-z0-9][A-Za-z0-9_]{1,31})\s*"
                r"(?P<assignment_syntax>:=|=)\s*(?P<default_value>.+?)\s*$"
            )
            structured_label = "selogic_equations"
        elif logic_mode:
            # Anchor at line start to avoid turning prose such as "setting ECOMM = N"
            # into an equation record. Names may start with digits (for example, 52A).
            assignment = re.compile(
                r"^\s*(?P<setting_name>[A-Za-z0-9][A-Za-z0-9_]{1,31})\s*=\s*"
                r"(?P<default_value>.+?)\s*$"
            )
            structured_label = "logic_equations"
        else:
            assignment = re.compile(
                r"(?P<setting_name>[A-Za-z][A-Za-z0-9_]{1,31})\s*:=\s*"
                r"(?P<default_value>.*?)(?=\s+[A-Za-z][A-Za-z0-9_]{1,31}\s*:=|$)"
            )
            structured_label = "setting_assignments"
        with pdfplumber.open(path) as pdf:
            for page_number in parse_number_spec(str(spec), len(pdf.pages)):
                text = pdf.pages[page_number - 1].extract_text(
                    layout=True,
                    x_density=float(pdf_cfg.get("x_density", 7.25)),
                    y_density=float(pdf_cfg.get("y_density", 13)),
                ) or ""
                raw_lines = text.splitlines()
                page_count = 0
                page_seen: dict[str, int] = {}
                for line_number, raw_line in enumerate(raw_lines, 1):
                    matches = list(assignment.finditer(raw_line))
                    for match_number, match in enumerate(matches, 1):
                        value = clean_value(match.group("default_value")) or ""
                        if not logic_mode and not programming_mode and match_number == len(matches):
                            value_col = match.start("default_value")
                            continuations = []
                            for following in raw_lines[line_number:]:
                                stripped = following.strip()
                                if not stripped or ":=" in following:
                                    break
                                indent = len(following) - len(following.lstrip())
                                if indent < max(0, value_col - 20):
                                    break
                                if re.match(r"(?:Table|Figure|Date Code|SEL-751|Protection and Logic|Logic Settings|Setting[s]? Prompt)", stripped):
                                    break
                                continuations.append(stripped)
                            if continuations:
                                value = " ".join([value, *continuations]).strip()
                        setting_name = clean_value(match.group("setting_name"))
                        if (logic_mode or programming_mode) and setting_name != setting_name.upper():
                            continue
                        if programming_mode and setting_name in {"LVALUE", "RVALUE"}:
                            continue
                        setting_key = normalize_key(setting_name)
                        page_seen[setting_key] = page_seen.get(setting_key, 0) + 1
                        if logic_mode or programming_mode:
                            syntax = match.group("assignment_syntax") if programming_mode else "="
                            data = {
                                "record_type": "selogic_equation" if programming_mode else "logic_equation",
                                "setting_name": setting_name,
                                "expression": value,
                                "assignment_syntax": syntax,
                                "dialect": "modern" if syntax == ":=" else "legacy",
                                "text": clean_value(raw_line),
                            }
                        else:
                            data = {
                                "record_type": "setting_assignment",
                                "setting_name": setting_name,
                                "default_value": value,
                                "assignment_syntax": ":=",
                                "text": clean_value(raw_line),
                            }
                        output.append(make_record(cfg["namespace"], path.name, "pdf",
                            {"page": page_number, "line": line_number,
                             "assignment": match_number, "occurrence": page_seen[setting_key],
                             "mode": structured_label}, data,
                            cfg.get("record_id_fields", [])))
                        page_count += 1
                line_count = 0
                if include_lines:
                    for line_number, raw_line in enumerate(raw_lines, 1):
                        text_line = clean_value(raw_line)
                        if not text_line:
                            continue
                        output.append(make_record(cfg["namespace"], path.name, "pdf",
                            {"page": page_number, "line": line_number, "mode": "layout_lines"},
                            {"record_type": "layout_line", "text": text_line},
                            cfg.get("record_id_fields", [])))
                        line_count += 1
                if page_count == 0 and not include_lines:
                    audit["warnings"].append(
                        f"No {structured_label.replace('_', ' ')} found on PDF page {page_number}; "
                        "verify layout or use a combined lines mode")
                selection_audit = {"page": page_number, "mode": mode,
                                   "records": page_count + line_count, "lines": line_count}
                selection_audit[structured_label] = page_count
                audit["processed_selections"].append(selection_audit)
        return output
    if mode != "lines":
        raise ValueError(f"Unsupported PDF mode: {mode}")
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise RuntimeError("PDF lines mode requires pypdf") from e
    reader = PdfReader(path)
    for page_number in parse_number_spec(str(spec), len(reader.pages)):
        text = reader.pages[page_number - 1].extract_text() or ""
        lines = [clean_value(x) for x in text.splitlines() if clean_value(x)]
        if not lines:
            audit["warnings"].append(f"No extractable text on PDF page {page_number}; OCR may be required")
        for line_number, text_line in enumerate(lines, 1):
            output.append(make_record(cfg["namespace"], path.name, "pdf",
                {"page": page_number, "line": line_number, "mode": mode},
                {"record_type": "line", "text": text_line},
                cfg.get("record_id_fields", [])))
        audit["processed_selections"].append({"page": page_number, "mode": mode, "records": len(lines)})
    return output


def build_indexes(records: list[dict[str, Any]], fields: list[str]) -> dict[str, dict[str, list[str]]]:
    indexes: dict[str, dict[str, list[str]]] = {field: {} for field in fields}
    for record in records:
        for field in fields:
            value = record["data"].get(field)
            if value in (None, ""):
                continue
            key = normalize_key(value)
            indexes[field].setdefault(key, []).append(record["record_id"])
    return indexes


def extract_records(source: Path, cfg: dict[str, Any], args: argparse.Namespace,
                    audit: dict[str, Any]) -> list[dict[str, Any]]:
    suffix = source.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return excel_records(source, cfg, args, audit)
    if suffix == ".csv":
        return csv_records(source, cfg, audit)
    if suffix == ".docx":
        return word_records(source, cfg, args, audit)
    if suffix == ".pdf":
        return pdf_records(source, cfg, args, audit)
    raise ValueError(f"Unsupported source type: {source.suffix}")


def deep_merge(base: dict[str, Any], override: dict[str, Any]) -> dict[str, Any]:
    result = dict(base)
    for key, value in override.items():
        if isinstance(value, dict) and isinstance(result.get(key), dict):
            result[key] = deep_merge(result[key], value)
        else:
            result[key] = value
    return result


def canonical_identity(namespace: str, field: str, value: Any) -> tuple[str, str]:
    key = normalize_key(value)
    identity = {"namespace": namespace, "field": field, "key": key}
    return hashlib.sha256(canonical_json(identity).encode()).hexdigest()[:24], key


def enrich_merged_records(records: list[dict[str, Any]], namespace: str, source_id: str,
                          metadata: dict[str, Any], merge_cfg: dict[str, Any]) -> None:
    canonical_field = merge_cfg.get("canonical_field")
    variant_fields = list(merge_cfg.get("variant_fields", []))
    for record in records:
        record["source"]["id"] = source_id
        record["source"]["metadata"] = {k: clean_value(v) for k, v in metadata.items()}
        for key, value in metadata.items():
            value = clean_value(value)
            if key in record["data"] and record["data"][key] != value:
                raise ValueError(f"Source metadata field collides with extracted data: {key}")
            record["data"][key] = value
        canonical_value = record["data"].get(canonical_field) if canonical_field else None
        if canonical_value not in (None, ""):
            canonical_id, canonical_key = canonical_identity(namespace, canonical_field, canonical_value)
            dimensions = {field: normalize_key(record["data"].get(field, ""))
                          for field in variant_fields}
            variant_id = hashlib.sha256(canonical_json({
                "canonical_id": canonical_id, "dimensions": dimensions
            }).encode()).hexdigest()[:24]
            record["data"]["canonical_id"] = canonical_id
            record["data"]["canonical_key"] = canonical_key
            record["data"]["variant_id"] = variant_id
            occurrence = {"variant_id": variant_id, "source_id": source_id,
                          "locator": record["locator"]}
            record["record_id"] = hashlib.sha256(
                canonical_json(occurrence).encode()).hexdigest()[:24]
        else:
            occurrence = {"namespace": namespace, "source_id": source_id,
                          "locator": record["locator"], "data_hash": record["content_hash"]}
            record["record_id"] = hashlib.sha256(
                canonical_json(occurrence).encode()).hexdigest()[:24]
        record["content_hash"] = hashlib.sha256(
            canonical_json(record["data"]).encode()).hexdigest()


def build_concepts(records: list[dict[str, Any]], canonical_field: str | None,
                   variant_fields: list[str]) -> list[dict[str, Any]]:
    if not canonical_field:
        return []
    concepts: dict[str, dict[str, Any]] = {}
    for record in records:
        data = record["data"]
        canonical_id = data.get("canonical_id")
        if not canonical_id:
            continue
        concept = concepts.setdefault(canonical_id, {
            "canonical_id": canonical_id,
            "canonical_field": canonical_field,
            "canonical_key": data["canonical_key"],
            "display_values": set(),
            "variants": {},
        })
        concept["display_values"].add(str(data[canonical_field]))
        variant_id = data["variant_id"]
        variant = concept["variants"].setdefault(variant_id, {
            "variant_id": variant_id,
            "dimensions": {field: data.get(field) for field in variant_fields},
            "record_ids": [],
        })
        variant["record_ids"].append(record["record_id"])
    output = []
    for canonical_id in sorted(concepts):
        concept = concepts[canonical_id]
        concept["display_values"] = sorted(concept["display_values"], key=normalize_key)
        concept["variants"] = [concept["variants"][key] for key in sorted(concept["variants"])]
        for variant in concept["variants"]:
            variant["record_ids"].sort()
        output.append(concept)
    return output


def safe_group_name(value: Any) -> str:
    name = re.sub(r"[^0-9a-z]+", "-", normalize_key(value)).strip("-")
    return name or "unclassified"


def write_dataset(records: list[dict[str, Any]], cfg: dict[str, Any], output: Path,
                  sources: list[dict[str, Any]], audit: dict[str, Any]) -> int:
    counts: dict[str, int] = {}
    for record in records:
        counts[record["record_id"]] = counts.get(record["record_id"], 0) + 1
    audit["duplicate_record_ids"] = sorted(k for k, v in counts.items() if v > 1)
    if audit["duplicate_record_ids"]:
        audit["errors"].append(
            "Duplicate record IDs detected; configure unique source IDs and variant fields")
    if not records:
        audit["errors"].append("Extraction produced zero records")
    output.mkdir(parents=True, exist_ok=True)
    for pattern in ("records-*.json", "concepts.json", "indexes.json", "manifest.json", "qaqc.json"):
        for old in output.glob(pattern):
            old.unlink()
    shard_size = max(1, int(cfg.get("shard_size", 5000)))
    shard_by = cfg.get("merge", {}).get("shard_by")
    grouped: dict[str, list[dict[str, Any]]] = {}
    if shard_by:
        for record in records:
            grouped.setdefault(safe_group_name(record["data"].get(shard_by)), []).append(record)
    else:
        grouped[""] = records
    shards, locations = [], {}
    for group in sorted(grouped):
        group_records = grouped[group]
        for start in range(0, len(group_records), shard_size):
            subset = group_records[start:start + shard_size]
            number = start // shard_size + 1
            filename = (f"records-{group}-{number:05d}.json" if group
                        else f"records-{number:05d}.json")
            payload = (canonical_json(subset) + "\n").encode("utf-8")
            (output / filename).write_bytes(payload)
            shard_hash = sha256_bytes(payload)
            descriptor = {"file": filename, "records": len(subset), "sha256": shard_hash}
            if group:
                descriptor["group"] = group
            shards.append(descriptor)
            for offset, record in enumerate(subset):
                locations[record["record_id"]] = {"shard": filename, "offset": offset}
    indexes = {"normalization": "NFKC + trim + collapse whitespace + casefold",
               "record_locations": locations,
               "fields": build_indexes(records, list(cfg.get("indexes", [])))}
    index_payload = (canonical_json(indexes) + "\n").encode("utf-8")
    (output / "indexes.json").write_bytes(index_payload)
    merge_cfg = cfg.get("merge", {})
    concepts = build_concepts(records, merge_cfg.get("canonical_field"),
                              list(merge_cfg.get("variant_fields", [])))
    concepts_payload = (canonical_json(concepts) + "\n").encode("utf-8")
    if concepts:
        (output / "concepts.json").write_bytes(concepts_payload)
    config_hash = sha256_bytes(canonical_json(cfg).encode())
    build_fingerprint = sha256_bytes(canonical_json({
        "schema": SCHEMA_VERSION,
        "sources": [{"id": s["id"], "sha256": s["sha256"]} for s in sources],
        "config": config_hash,
        "shards": shards,
        "concepts": sha256_bytes(concepts_payload) if concepts else None,
    }).encode())
    manifest = {"schema_version": SCHEMA_VERSION, "dataset_name": cfg["dataset_name"],
                "namespace": cfg["namespace"], "sources": sources,
                "config_sha256": config_hash, "build_fingerprint": build_fingerprint,
                "record_count": len(records), "concept_count": len(concepts),
                "shards": shards, "shard_by": shard_by, "index_file": "indexes.json",
                "concept_file": "concepts.json" if concepts else None,
                "index_fields": list(cfg.get("indexes", [])), "warnings": audit["warnings"]}
    if len(sources) == 1:
        manifest["source"] = {"name": sources[0]["name"], "sha256": sources[0]["sha256"]}
    manifest_payload = (canonical_json(manifest) + "\n").encode("utf-8")
    (output / "manifest.json").write_bytes(manifest_payload)
    audit["record_count"] = len(records)
    audit["concept_count"] = len(concepts)
    audit["output_hashes"] = {"manifest.json": sha256_bytes(manifest_payload),
                              "indexes.json": sha256_bytes(index_payload),
                              **({"concepts.json": sha256_bytes(concepts_payload)} if concepts else {}),
                              **{s["file"]: s["sha256"] for s in shards}}
    qaqc_payload = (canonical_json(audit) + "\n").encode("utf-8")
    (output / "qaqc.json").write_bytes(qaqc_payload)
    print(json.dumps({"output": str(output), "sources": len(sources),
                      "records": len(records), "concepts": len(concepts),
                      "shards": len(shards), "warnings": audit["warnings"],
                      "errors": audit["errors"]}, indent=2))
    return 2 if audit["errors"] else 0


def default_args() -> argparse.Namespace:
    return argparse.Namespace(pages=None, sheets=[], ranges=[], word_tables=[])


def build(args: argparse.Namespace) -> int:
    source, output = Path(args.source).resolve(), Path(args.output).resolve()
    if not source.is_file():
        raise FileNotFoundError(source)
    if source.suffix.lower() not in SUPPORTED:
        raise ValueError(f"Unsupported source type: {source.suffix}")
    cfg = json.loads(Path(args.config).read_text(encoding="utf-8"))
    cfg.setdefault("dataset_name", source.stem)
    cfg.setdefault("namespace", cfg["dataset_name"])
    cfg.setdefault("record_id_fields", [])
    cfg.setdefault("indexes", [])
    cfg.setdefault("shard_size", 5000)
    audit = {"requested_source": str(source), "processed_selections": [], "warnings": [],
             "errors": [], "duplicate_record_ids": [], "record_count": 0, "output_hashes": {}}
    records = extract_records(source, cfg, args, audit)
    source_info = {"id": source.stem, "name": source.name,
                   "type": source.suffix.lower().lstrip("."), "sha256": sha256_file(source),
                   "metadata": {}}
    return write_dataset(records, cfg, output, [source_info], audit)


def parse_source_overrides(values: list[str]) -> dict[str, Path]:
    overrides: dict[str, Path] = {}
    for value in values:
        if "=" not in value:
            raise ValueError(f"Source override must be ID=PATH: {value}")
        source_id, path = value.split("=", 1)
        if not source_id.strip() or not path.strip():
            raise ValueError(f"Source override must be ID=PATH: {value}")
        overrides[source_id.strip()] = Path(path).expanduser().resolve()
    return overrides


def build_multi(args: argparse.Namespace) -> int:
    config_path = Path(args.config).resolve()
    cfg = json.loads(config_path.read_text(encoding="utf-8"))
    cfg.setdefault("dataset_name", "merged-dataset")
    cfg.setdefault("namespace", cfg["dataset_name"])
    cfg.setdefault("record_id_fields", [])
    cfg.setdefault("indexes", [])
    cfg.setdefault("shard_size", 5000)
    source_items = cfg.get("sources", [])
    if not source_items:
        raise ValueError("Multi-source config requires a nonempty sources array")
    overrides = parse_source_overrides(args.source_overrides)
    known_ids = {str(item.get("id", "")) for item in source_items}
    unknown = sorted(set(overrides) - known_ids)
    if unknown:
        raise ValueError(f"Source override IDs not present in config: {unknown}")
    audit = {"requested_sources": [], "processed_sources": [], "processed_selections": [],
             "warnings": [], "errors": [], "duplicate_record_ids": [], "record_count": 0,
             "output_hashes": {}}
    all_records: list[dict[str, Any]] = []
    source_infos: list[dict[str, Any]] = []
    seen_ids: set[str] = set()
    global_base = {key: value for key, value in cfg.items()
                   if key not in {"sources", "merge"}}
    for item in source_items:
        source_id = str(item.get("id", "")).strip()
        if not source_id:
            raise ValueError("Every source requires a stable id")
        if source_id in seen_ids:
            raise ValueError(f"Duplicate source id: {source_id}")
        seen_ids.add(source_id)
        configured_path = item.get("path")
        if source_id in overrides:
            source = overrides[source_id]
        elif configured_path:
            source = Path(configured_path).expanduser()
            if not source.is_absolute():
                source = (config_path.parent / source).resolve()
        else:
            raise ValueError(f"Source {source_id} requires path or --source {source_id}=PATH")
        if not source.is_file():
            raise FileNotFoundError(source)
        if source.suffix.lower() not in SUPPORTED:
            raise ValueError(f"Unsupported source type: {source.suffix}")
        item_cfg = {key: value for key, value in item.items()
                    if key not in {"id", "path", "metadata"}}
        source_cfg = deep_merge(global_base, item_cfg)
        source_cfg["namespace"] = cfg["namespace"]
        source_audit = {"processed_selections": [], "warnings": [], "errors": []}
        records = extract_records(source, source_cfg, default_args(), source_audit)
        metadata = dict(item.get("metadata", {}))
        enrich_merged_records(records, cfg["namespace"], source_id, metadata,
                              cfg.get("merge", {}))
        all_records.extend(records)
        audit["requested_sources"].append({"id": source_id, "path": str(source)})
        audit["processed_sources"].append({"id": source_id, "name": source.name,
                                           "records": len(records),
                                           "warnings": source_audit["warnings"],
                                           "errors": source_audit["errors"]})
        audit["processed_selections"].extend(
            [{"source_id": source_id, **selection}
             for selection in source_audit["processed_selections"]])
        audit["warnings"].extend(f"{source_id}: {x}" for x in source_audit["warnings"])
        audit["errors"].extend(f"{source_id}: {x}" for x in source_audit["errors"])
        source_infos.append({"id": source_id, "name": source.name,
                             "type": source.suffix.lower().lstrip("."),
                             "sha256": sha256_file(source), "metadata": metadata})
    return write_dataset(all_records, cfg, Path(args.output).resolve(), source_infos, audit)


def parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description=__doc__)
    sub = p.add_subparsers(dest="command", required=True)
    b = sub.add_parser("build", help="Build a lookup dataset from one source")
    b.add_argument("--source", required=True)
    b.add_argument("--config", required=True)
    b.add_argument("--output", required=True)
    b.add_argument("--sheet", dest="sheets", action="append", default=[])
    b.add_argument("--range", dest="ranges", action="append", default=[])
    b.add_argument("--page", dest="pages", help="PDF pages, e.g. 1-5,8")
    b.add_argument("--word-table", dest="word_tables", action="append", type=int, default=[])
    m = sub.add_parser("build-multi", help="Build one normalized dataset from multiple sources")
    m.add_argument("--config", required=True)
    m.add_argument("--output", required=True)
    m.add_argument("--source", dest="source_overrides", action="append", default=[],
                   metavar="ID=PATH", help="Override one configured source path")
    return p


if __name__ == "__main__":
    arguments = parser().parse_args()
    try:
        if arguments.command == "build":
            sys.exit(build(arguments))
        if arguments.command == "build-multi":
            sys.exit(build_multi(arguments))
        sys.exit(1)
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        sys.exit(1)
