#!/usr/bin/env python3
import json, subprocess, sys, tempfile, unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
BUILDER = ROOT / "scripts" / "document_to_json.py"
LOOKUP = ROOT / "scripts" / "lookup_json.py"


class FrameworkTests(unittest.TestCase):
    def run_build(self, source: Path, config: dict, out: Path):
        cfg = out.parent / "config.json"
        cfg.write_text(json.dumps(config), encoding="utf-8")
        return subprocess.run([sys.executable, str(BUILDER), "build", "--source", str(source),
                               "--config", str(cfg), "--output", str(out)],
                              text=True, capture_output=True)

    def base_config(self):
        return {"dataset_name": "test", "namespace": "test-v1", "record_id_fields": ["tag"],
                "indexes": ["tag"], "shard_size": 1, "selection": {}}

    def run_multi_build(self, config: dict, out: Path):
        cfg = out.parent / f"{out.name}-multi-config.json"
        cfg.write_text(json.dumps(config), encoding="utf-8")
        return subprocess.run([sys.executable, str(BUILDER), "build-multi",
                               "--config", str(cfg), "--output", str(out)],
                              text=True, capture_output=True)

    def test_csv_build_repeatability_and_lookup(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); source = root / "data.csv"
            source.write_text("Tag,Description\n52A,Breaker\nPT-1,Potential Transformer\n", encoding="utf-8")
            cfg = self.base_config(); cfg["selection"] = {"csv": {"header_row": 1}}
            out1, out2 = root / "out1", root / "out2"
            self.assertEqual(self.run_build(source, cfg, out1).returncode, 0)
            self.assertEqual(self.run_build(source, cfg, out2).returncode, 0)
            m1 = json.loads((out1 / "manifest.json").read_text())
            m2 = json.loads((out2 / "manifest.json").read_text())
            self.assertEqual(m1["build_fingerprint"], m2["build_fingerprint"])
            result = subprocess.run([sys.executable, str(LOOKUP), "--dataset", str(out1),
                                     "--index", "tag", "--key", " 52a "],
                                    text=True, capture_output=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual(json.loads(result.stdout)["records"][0]["data"]["description"], "Breaker")

    def test_excel_sheet_range(self):
        try:
            from openpyxl import Workbook
        except ImportError:
            self.skipTest("openpyxl unavailable")
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); source = root / "data.xlsx"
            wb = Workbook(); ws = wb.active; ws.title = "Equipment"
            ws.append(["Tag", "Description"]); ws.append(["52A", "Breaker"]); ws.append(["52B", "Spare"])
            wb.create_sheet("Ignored").append(["X"]); wb.save(source)
            cfg = self.base_config(); cfg["selection"] = {"spreadsheets": {"sheets": [
                {"name": "Equipment", "range": "A1:B2", "header_row": 1}]}}
            out = root / "out"; result = self.run_build(source, cfg, out)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual(json.loads((out / "manifest.json").read_text())["record_count"], 1)

    def test_word_table_and_paragraph(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx unavailable")
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); source = root / "data.docx"
            doc = Document(); doc.add_paragraph("Reference heading")
            table = doc.add_table(rows=2, cols=2); table.cell(0,0).text="Tag"; table.cell(0,1).text="Description"
            table.cell(1,0).text="52A"; table.cell(1,1).text="Breaker"; doc.save(source)
            cfg = self.base_config(); cfg["selection"] = {"word": {"paragraphs": "1", "tables": [
                {"index": 0, "header_row": 1}]}}
            out = root / "out"; result = self.run_build(source, cfg, out)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual(json.loads((out / "manifest.json").read_text())["record_count"], 2)

    def test_pdf_page_text(self):
        try:
            from reportlab.pdfgen import canvas
            import pypdf  # noqa: F401
        except ImportError:
            self.skipTest("reportlab or pypdf unavailable")
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); source = root / "data.pdf"
            c = canvas.Canvas(str(source)); c.drawString(72, 720, "SELOGIC LATCHES N, 1-64 ELAT := 4"); c.save()
            cfg = self.base_config(); cfg["record_id_fields"] = ["setting_name"]
            cfg["indexes"] = ["setting_name"]
            cfg["selection"] = {"pdf": {"pages": "1", "mode": "setting_assignments"}}
            out = root / "out"; result = self.run_build(source, cfg, out)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual(json.loads((out / "manifest.json").read_text())["record_count"], 1)
            lookup = subprocess.run([sys.executable, str(LOOKUP), "--dataset", str(out),
                                     "--index", "setting_name", "--key", "elat"],
                                    text=True, capture_output=True)
            self.assertEqual(lookup.returncode, 0, lookup.stderr)
            record = json.loads(lookup.stdout)["records"][0]
            self.assertEqual(record["data"]["default_value"], "4")


    def test_pdf_logic_equations(self):
        try:
            from reportlab.pdfgen import canvas
            import pdfplumber  # noqa: F401
        except ImportError:
            self.skipTest("reportlab or pdfplumber unavailable")
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); source = root / "logic.pdf"
            c = canvas.Canvas(str(source))
            c.drawString(72, 720, "ER = /51P1 + /51G1 + /OUT103")
            c.drawString(72, 700, "52A = !IN101")
            c.drawString(72, 680, "If setting ECOMM = N, this is prose.")
            c.drawString(72, 660, "shot = 0. This prose begins with a lowercase word.")
            c.save()
            cfg = self.base_config(); cfg["record_id_fields"] = ["setting_name"]
            cfg["indexes"] = ["setting_name"]
            cfg["selection"] = {"pdf": {"pages": "1", "mode": "logic_equations"}}
            out = root / "out"; result = self.run_build(source, cfg, out)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual(json.loads((out / "manifest.json").read_text())["record_count"], 2)
            lookup = subprocess.run([sys.executable, str(LOOKUP), "--dataset", str(out),
                                     "--index", "setting_name", "--key", "52a"],
                                    text=True, capture_output=True)
            self.assertEqual(lookup.returncode, 0, lookup.stderr)
            record = json.loads(lookup.stdout)["records"][0]
            self.assertEqual(record["data"]["expression"], "!IN101")
            self.assertEqual(record["data"]["record_type"], "logic_equation")


    def test_pdf_selogic_programming_mixed_syntax(self):
        try:
            from reportlab.pdfgen import canvas
            import pdfplumber  # noqa: F401
        except ImportError:
            self.skipTest("reportlab or pdfplumber unavailable")
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); source = root / "programming.pdf"
            c = canvas.Canvas(str(source))
            c.drawString(72, 720, "PMV01 := 5 # Modern assignment")
            c.drawString(72, 700, "SV1 = IN101 + RB3 * LT4")
            c.drawString(72, 680, "LVALUE := Expression")
            c.drawString(72, 660, "If setting ECOMM = N, this is prose.")
            c.drawString(72, 640, "shot = 0. This prose begins with a lowercase word.")
            c.save()
            cfg = self.base_config(); cfg["record_id_fields"] = ["setting_name"]
            cfg["indexes"] = ["setting_name", "assignment_syntax"]
            cfg["selection"] = {"pdf": {"pages": "1", "mode": "selogic_programming"}}
            out = root / "out"; result = self.run_build(source, cfg, out)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual(json.loads((out / "manifest.json").read_text())["record_count"], 2)
            lookup = subprocess.run([sys.executable, str(LOOKUP), "--dataset", str(out),
                                     "--index", "setting_name", "--key", "pmv01"],
                                    text=True, capture_output=True)
            self.assertEqual(lookup.returncode, 0, lookup.stderr)
            record = json.loads(lookup.stdout)["records"][0]
            self.assertEqual(record["data"]["expression"], "5 # Modern assignment")
            self.assertEqual(record["data"]["assignment_syntax"], ":=")
            self.assertEqual(record["data"]["dialect"], "modern")

    def test_multi_source_canonical_variants_and_filtered_lookup(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            first, second = root / "first.csv", root / "second.csv"
            first.write_text("Tag,Expression\nOUT101,IN101\n", encoding="utf-8")
            second.write_text("Tag,Expression\nOUT101,IN201\n", encoding="utf-8")
            cfg = {
                "dataset_name": "merged-test",
                "namespace": "merged-test-v1",
                "record_id_fields": ["tag"],
                "indexes": ["tag", "product_family", "canonical_id", "variant_id"],
                "shard_size": 100,
                "merge": {
                    "canonical_field": "tag",
                    "variant_fields": ["product_family", "product_model"],
                    "shard_by": "product_family"
                },
                "sources": [
                    {"id": "family-3xx", "path": str(first),
                     "metadata": {"product_family": "3xx", "product_model": "SEL-351S"},
                     "selection": {"csv": {"header_row": 1}}},
                    {"id": "family-7xx", "path": str(second),
                     "metadata": {"product_family": "7xx", "product_model": "SEL-751"},
                     "selection": {"csv": {"header_row": 1}}}
                ]
            }
            out1, out2 = root / "merged1", root / "merged2"
            result1 = self.run_multi_build(cfg, out1)
            result2 = self.run_multi_build(cfg, out2)
            self.assertEqual(result1.returncode, 0, result1.stderr)
            self.assertEqual(result2.returncode, 0, result2.stderr)
            manifest1 = json.loads((out1 / "manifest.json").read_text())
            manifest2 = json.loads((out2 / "manifest.json").read_text())
            self.assertEqual(manifest1["record_count"], 2)
            self.assertEqual(manifest1["concept_count"], 1)
            self.assertEqual(manifest1["build_fingerprint"], manifest2["build_fingerprint"])
            concepts = json.loads((out1 / "concepts.json").read_text())
            self.assertEqual(len(concepts[0]["variants"]), 2)
            lookup = subprocess.run([sys.executable, str(LOOKUP), "--dataset", str(out1),
                                     "--index", "tag", "--key", "out101",
                                     "--where", "product_family=7XX"],
                                    text=True, capture_output=True)
            self.assertEqual(lookup.returncode, 0, lookup.stderr)
            payload = json.loads(lookup.stdout)
            self.assertEqual(payload["matches"], 1)
            self.assertEqual(payload["records"][0]["data"]["expression"], "IN201")
            self.assertEqual(payload["records"][0]["source"]["id"], "family-7xx")


if __name__ == "__main__":
    unittest.main()
